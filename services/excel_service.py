"""Talabalar ro'yxatini (Excel) o'qish va talabani qidirish.

HEMIS ID faqat Excel fayldagi qiymatdan olinadi: generatsiya qilinmaydi,
taxmin qilinmaydi va o'zgartirilmaydi.
"""
import logging
import os
import re
import threading
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from config import EXCEL_FILE
from services.validation_service import (
    normalize_full_name,
    normalize_passport,
    normalize_phone,
)

logger = logging.getLogger(__name__)

# Sarlavha qatori shu qatorlar ichidan qidiriladi
HEADER_SCAN_ROWS = 10


class ExcelDataError(Exception):
    """Excel fayl topilmadi yoki undagi ustunlarni aniqlab bo'lmadi."""


@dataclass(frozen=True)
class Student:
    full_name: str        # standartlashtirilgan F.I.Sh.
    phones: frozenset     # 998XXXXXXXXX ko'rinishidagi raqamlar
    passport: str         # standartlashtirilgan pasport
    hemis_id: str         # Excel'dagi qiymat (o'zgartirilmagan)


@dataclass(frozen=True)
class Columns:
    hemis: int
    passport: int | None
    phone: int | None
    full_name: int | None
    surname: int | None
    first_name: int | None
    patronymic: int | None


def _header_key(value) -> str:
    """Sarlavhani solishtirish uchun: kichik harf, faqat harf va raqamlar."""
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def find_hemis_id_column(headers: list) -> int | None:
    """'HEMIS ID', 'HEMIS_ID', 'HEMISID', 'Hemis ID' kabi ustunni topadi."""
    keys = [_header_key(h) for h in headers]
    for exact in ("hemisid", "hemis"):
        if exact in keys:
            return keys.index(exact)
    for i, key in enumerate(keys):
        if "hemis" in key:
            return i
    return None


def find_passport_column(headers: list) -> int | None:
    """'Pasport', 'Passport', 'Pasport seriya va raqami' kabi ustunni topadi."""
    for i, h in enumerate(headers):
        key = _header_key(h)
        if "pasport" in key or "passport" in key:
            return i
    return None


def find_phone_column(headers: list) -> int | None:
    """'Telefon', 'Telefon raqam', 'Phone', 'Telefon №' kabi ustunni topadi."""
    for i, h in enumerate(headers):
        key = _header_key(h)
        if "telefon" in key or "phone" in key or key in ("tel", "telraqam", "mobil"):
            return i
    return None


def find_name_columns(headers: list) -> dict:
    """F.I.Sh. ustunini (yoki alohida Familiya / Ism / Otasining ismi ustunlarini) topadi."""
    result = {"full_name": None, "surname": None, "first_name": None, "patronymic": None}
    for i, h in enumerate(headers):
        key = _header_key(h)
        if not key:
            continue
        if result["full_name"] is None and (
            key.startswith(("fio", "fish", "fullname"))
            or key in ("ismfamiliya", "familiyaism", "familiyaismi", "talabafio",
                       "talabafish", "talaba", "name")
        ):
            result["full_name"] = i
        elif key in ("familiya", "familiyasi", "surname", "lastname"):
            result["surname"] = i
        elif key in ("ism", "ismi", "firstname"):
            result["first_name"] = i
        elif key in ("otasiningismi", "sharifi", "patronymic", "middlename"):
            result["patronymic"] = i
    return result


def detect_columns(headers: list) -> Columns | None:
    hemis = find_hemis_id_column(headers)
    if hemis is None:
        return None
    names = find_name_columns(headers)
    return Columns(
        hemis=hemis,
        passport=find_passport_column(headers),
        phone=find_phone_column(headers),
        **names,
    )


def _cell_text(value) -> str:
    """Excel katagini matnga aynan aylantiradi (butun son 12345.0 -> '12345')."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _row_phones(value) -> frozenset:
    """Katakda bir nechta raqam bo'lishi mumkin (vergul, nuqtali vergul, / bilan)."""
    phones = set()
    for part in re.split(r"[,;/\n]", _cell_text(value)):
        phone = normalize_phone(part)
        if phone:
            phones.add(phone)
    return frozenset(phones)


def _row_name(row: tuple, cols: Columns) -> str:
    def get(idx):
        return _cell_text(row[idx]) if idx is not None and idx < len(row) else ""

    if cols.full_name is not None:
        return normalize_full_name(get(cols.full_name))
    parts = [get(cols.surname), get(cols.first_name), get(cols.patronymic)]
    return normalize_full_name(" ".join(p for p in parts if p))


def load_students(path: Path) -> list[Student]:
    """Excel fayldagi barcha varaqlardan talabalarni o'qiydi."""
    if not path.exists():
        raise ExcelDataError(f"Excel fayl topilmadi: {path.name}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    students: list[Student] = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            cols = None
            for _ in range(HEADER_SCAN_ROWS):
                header = next(rows, None)
                if header is None:
                    break
                cols = detect_columns(list(header))
                if cols:
                    break
            if not cols:
                continue

            for row in rows:
                if not row or cols.hemis >= len(row):
                    continue
                hemis_id = _cell_text(row[cols.hemis])
                if not hemis_id:
                    continue
                passport = ""
                if cols.passport is not None and cols.passport < len(row):
                    passport = normalize_passport(_cell_text(row[cols.passport]))
                phones = frozenset()
                if cols.phone is not None and cols.phone < len(row):
                    phones = _row_phones(row[cols.phone])
                students.append(Student(
                    full_name=_row_name(row, cols),
                    phones=phones,
                    passport=passport,
                    hemis_id=hemis_id,
                ))
    finally:
        workbook.close()

    if not students:
        raise ExcelDataError("Excel faylda HEMIS ID ustuni yoki talabalar topilmadi")
    return students


class StudentRegistry:
    """Excel ma'lumotlarini xotirada saqlaydi; fayl o'zgarsa avtomatik qayta o'qiydi."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self._lock = threading.Lock()
        self._mtime = None
        self._students: list[Student] = []

    def students(self) -> list[Student]:
        with self._lock:
            try:
                mtime = os.path.getmtime(self.path)
            except OSError:
                raise ExcelDataError(f"Excel fayl topilmadi: {self.path.name}")
            if mtime != self._mtime:
                self._students = load_students(self.path)
                self._mtime = mtime
                logger.info("Excel yuklandi: %d ta talaba", len(self._students))
            return self._students

    def find_student(self, full_name: str | None = None, phone: str | None = None,
                     passport: str | None = None) -> Student | None:
        return find_student(self.students(), full_name, phone, passport)


def _name_matches(entered: str, stored: str) -> bool:
    """Kiritilgan so'zlar Excel'dagi F.I.Sh. boshidagi so'zlar bilan bir xil bo'lishi kerak.

    Masalan, 'FAMILIYA ISM' Excel'dagi 'FAMILIYA ISM OTASINING_ISMI QIZI' ga mos keladi.
    """
    entered_words = entered.split()
    stored_words = stored.split()
    return len(entered_words) <= len(stored_words) and stored_words[:len(entered_words)] == entered_words


def find_student(students: list[Student], full_name: str | None = None,
                 phone: str | None = None, passport: str | None = None) -> Student | None:
    """Kiritilgan ma'lumotlarga to'liq mos keladigan yagona talabani qaytaradi.

    Qoidalar:
    - Excel'da mavjud bo'lgan har bir solishtiriladigan maydon mos kelishi shart;
      birortasi farq qilsa, bu qator rad etiladi.
    - Kamida 2 ta maydon mos kelishi va ulardan biri pasport yoki telefon bo'lishi shart
      (faqat ism-familiya bo'yicha HEMIS ID berilmaydi).
    - Bir nechta turli HEMIS ID mos kelsa, hech biri qaytarilmaydi.
    """
    name = normalize_full_name(full_name) if full_name else ""
    phone = normalize_phone(phone) if phone else None
    passport = normalize_passport(passport) if passport else ""

    matches: dict[str, Student] = {}
    for student in students:
        matched = set()

        if passport and student.passport:
            if student.passport != passport:
                continue
            matched.add("passport")
        if phone and student.phones:
            if phone not in student.phones:
                continue
            matched.add("phone")
        if name and student.full_name:
            if not _name_matches(name, student.full_name):
                continue
            matched.add("name")

        if len(matched) >= 2 and matched & {"passport", "phone"}:
            matches.setdefault(student.hemis_id, student)

    if len(matches) == 1:
        return next(iter(matches.values()))
    if len(matches) > 1:
        logger.warning("Bir nechta talaba mos keldi (%d ta), HEMIS ID berilmadi", len(matches))
    return None


registry = StudentRegistry(EXCEL_FILE)
